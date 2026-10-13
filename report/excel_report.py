from __future__ import annotations

import csv
from pathlib import Path
from typing import Iterable

from openpyxl import Workbook
from openpyxl.formatting.rule import ColorScaleRule
from openpyxl.styles import Alignment, Font, PatternFill
from openpyxl.utils import get_column_letter

from audit_types import AuditRecord


def _summary_row(record: AuditRecord) -> list:
    issues = record.issues
    critical = sum(1 for i in issues if i.level == "critical")
    warnings = sum(1 for i in issues if i.level == "warning")
    top3 = " | ".join(i.title for i in issues[:3])
    lead = record.lead
    tech = record.checks.get("technology", {})
    s = record.scores
    return [
        lead.lead_id, lead.company, lead.website_url, lead.contact_name, lead.email,
        tech.get("platform", "Unknown"), s.get("overall"), s.get("performance"),
        s.get("seo"), s.get("security"), s.get("ux_mobile"), s.get("ecommerce"),
        critical, warnings, top3, s.get("coverage_percent"), round(record.duration_seconds, 1),
    ]


def write_summary_csv(records: Iterable[AuditRecord], path: str | Path) -> None:
    path = Path(path); path.parent.mkdir(parents=True, exist_ok=True)
    headers = ["lead_id","company","website","contact_name","email","platform","overall_score","performance_score","seo_score","security_score","ux_mobile_score","ecommerce_score","critical_issues","warning_issues","top_3_issues","audit_coverage_percent","duration_seconds"]
    with path.open("w", encoding="utf-8-sig", newline="") as f:
        w = csv.writer(f); w.writerow(headers)
        for rec in records: w.writerow(_summary_row(rec))


def write_workbook(records: list[AuditRecord], path: str | Path) -> None:
    path = Path(path); path.parent.mkdir(parents=True, exist_ok=True)
    wb = Workbook()
    ws = wb.active; ws.title = "Summary"
    headers = ["Lead ID","Company","Website","Contact","Email","Platform","Overall","Performance","SEO","Security","UX/Mobile","E-commerce","Critical","Warnings","Top 3 Issues","Audit Coverage %","Seconds"]
    ws.append(headers)
    for rec in records: ws.append(_summary_row(rec))
    _style_header(ws)
    ws.freeze_panes = "A2"; ws.auto_filter.ref = ws.dimensions
    for col in range(7, 13):
        letter = get_column_letter(col)
        ws.conditional_formatting.add(f"{letter}2:{letter}{max(2, ws.max_row)}", ColorScaleRule(start_type='num', start_value=0, start_color='F8696B', mid_type='num', mid_value=70, mid_color='FFEB84', end_type='num', end_value=100, end_color='63BE7B'))
    _autosize(ws)

    issues_ws = wb.create_sheet("All Issues")
    issues_ws.append(["Lead ID","Company","Website","Category","Level","Points","Issue","Evidence","How to Fix"])
    for rec in records:
        for issue in rec.issues:
            issues_ws.append([rec.lead.lead_id, rec.lead.company, rec.lead.website_url, issue.category, issue.level, issue.points, issue.title, issue.evidence, issue.fix])
    _style_header(issues_ws); issues_ws.freeze_panes = "A2"; issues_ws.auto_filter.ref = issues_ws.dimensions; _autosize(issues_ws, max_width=60)

    raw_ws = wb.create_sheet("Raw Metrics")
    raw_ws.append(["Lead ID","Company","Check","Metric","Value"])
    for rec in records:
        for check_name, check_data in rec.checks.items():
            for key, value in _flatten(check_data):
                raw_ws.append([rec.lead.lead_id, rec.lead.company, check_name, key, str(value)[:32000]])
    _style_header(raw_ws); raw_ws.freeze_panes = "A2"; raw_ws.auto_filter.ref = raw_ws.dimensions; _autosize(raw_ws, max_width=70)
    wb.save(path)


def _flatten(value, prefix=""):
    if isinstance(value, dict):
        for k, v in value.items():
            p = f"{prefix}.{k}" if prefix else str(k)
            yield from _flatten(v, p)
    elif isinstance(value, list):
        yield prefix, value
    else:
        yield prefix, value


def _style_header(ws):
    fill = PatternFill("solid", fgColor="1F4E78")
    for cell in ws[1]:
        cell.font = Font(color="FFFFFF", bold=True); cell.fill = fill; cell.alignment = Alignment(vertical="center")


def _autosize(ws, max_width=45):
    for col_cells in ws.columns:
        width = 10
        for cell in col_cells[: min(len(col_cells), 200)]:
            if cell.value is not None:
                width = max(width, min(max_width, len(str(cell.value)) + 2))
        ws.column_dimensions[get_column_letter(col_cells[0].column)].width = width
